"""Read one source file into Units - every format except the textbook.

Each format has its own idea of structure: HTML headings, JATS <sec>
elements, Word heading styles, Markdown #'s, and for a PDF nothing but font
sizes and bold. Each loader here reduces its format to the same thing,
chunking.Units in reading order, each tagged with

  section  the heading path, outermost first:
           ("Skin Cancer Treatment (PDQ)", "Treatment of BCC", "Mohs surgery")
  locator  where a citation points: "p. 12" for PDFs, else the innermost
           heading (the document title before the first heading)
  page     1-based page number for PDFs, None for everything else
  kind     prose | caption | reference | table

chunking.pack() then turns Units into chunks. The one size decision made
here is for tables: pack() never splits a non-prose unit, so a table is cut
into row groups of at most config.CHUNK_TARGET_CHARS, each repeating the
caption and header row, so every piece still says what its columns mean.
Rows are joined with newlines, cells with " | ".

The textbook (format pdf_textbook) is not read here: ingest.py keeps its own
block-based path so its chunk IDs and vectors stay exactly as they are.

Options (sources.yaml `options:`; any other key is an error)
  skip_tables       true -> drop tables, captions included. For documents
                    whose tables bury the prose, e.g. the ICD-10 code lists
                    in a Medicare billing article. PDFs have no table
                    detection, so it changes nothing there.
  exclude_sections  heading names whose content is dropped, subsections
                    included. Case-insensitive; a leading "1." / "F." and a
                    trailing "(05/12/2025)" or ":" on the heading are
                    ignored. REPLACES the defaults below, so
                    `exclude_sections: []` keeps everything.
                    In PDFs, where heading levels are guessed from font
                    sizes, only the text up to the next heading (of any
                    level) is dropped: a "Contents" heading ranked too high
                    must not swallow the rest of the document. The same
                    goes for what counts as "under References".

Default exclude_sections, and why each goes:
  StatPearls (NCBI Bookshelf)
    Continuing Education Activity - the article restated as a summary plus
        learning objectives ("Identify...", "Describe..."). Dense and on
        topic, so it would outrank the passages that actually state facts.
    Review Questions, Comment on this article - quiz and feedback links.
  NCI PDQ
    About This PDQ Summary - the same boilerplate on every PDQ page.
    Latest Updates to This Summary, Changes to This Summary - a changelog
        ("Added text to state...") that reads like current guidance.
    Current Clinical Trials - a pointer to the trial search, repeated per
        section.
  any document
    Table of Contents, Contents - headings and page numbers, no content.
    Acknowledgments, Funding, Conflicts of interest, Disclosure(s), Author
        contributions, Author information - back matter, never an answer.

HTML page furniture ("In This Section", "Similar articles", ...) is dropped
always, whatever the options say - it is never content.

Why "references" is a kind and not a deletion: everything under a
References / Bibliography heading becomes one reference unit per entry.
Retrieval leaves them out by default (config.RETRIEVE_KINDS), but they stay
in the index for anyone who wants to look a citation up.

Tables, unlike references, are searched like prose (RETRIEVE_KINDS has
"table"); skip_tables is how a source keeps its tables out of the index.

Duplicates: ingest refuses two chunks with one ID. An ID hashes the text
and its place: the page for PDFs, the section path for everything else. So
a unit is kept once per place: text repeated verbatim in one section, or
on one PDF page, is a duplicate; the same text under a same-named
subsection of another section ("Cryosurgery" under both BCC and SCC) is
not, and both stay. References are the exception: one repeated in two
sections' "References" lists is kept once, under the first.

Third-party parsers (BeautifulSoup, python-docx, PyMuPDF) are imported
inside their loader: only ingestion needs them, never the app.
"""

import re
import xml.etree.ElementTree as ET
from collections import Counter
from html.entities import name2codepoint
from math import ceil
from pathlib import Path

from src import config
from src.chunking import Unit, clean, split_long

# Same pattern as ingest.CAPTION_RE (a test keeps them equal). Copied, not
# imported: ingest imports this module, and the import must not go in a circle.
CAPTION_RE = re.compile(r"^(Fig\.?|Figure|Table|Video)\b", re.IGNORECASE)

DEFAULT_EXCLUDE_SECTIONS = (
    # StatPearls / NCBI Bookshelf
    "Continuing Education Activity",
    "Review Questions",
    "Comment on this article",
    # NCI PDQ
    "About This PDQ Summary",
    "Latest Updates to This Summary",
    "Changes to This Summary",
    "Current Clinical Trials",
    # any document
    "Table of Contents",
    "Contents",
    "Acknowledgments",
    "Acknowledgements",
    "Acknowledgment",
    "Funding",
    "Conflicts of Interest",
    "Conflict of Interest",
    "Disclosure",
    "Disclosures",
    "Author Contributions",
    "Author Information",
    "Author Information and Affiliations",
)

REFERENCE_HEADINGS = frozenset({
    "references", "reference", "bibliography", "reference list", "works cited",
    "literature cited", "references and notes", "further reading",
    "suggested reading", "suggested readings",
})

OPTIONS = {"skip_tables": "true/false", "exclude_sections": "a list of heading names"}


# ============================================================== options
def parse_options(options, source_id="source"):
    """sources.yaml `options:` -> {"skip_tables": bool, "exclude_sections": tuple}.

    Strict on purpose: a typo like `skip_table: true` would otherwise be
    ignored without a word, and the tables would be indexed anyway.
    """
    options = dict(options or {})
    unknown = sorted(set(options) - set(OPTIONS))
    if unknown:
        allowed = ", ".join(f"{k} ({v})" for k, v in OPTIONS.items())
        raise ValueError(f"{source_id}: unknown loader option(s) {unknown}. "
                         f"Allowed: {allowed}")
    skip = options.get("skip_tables", False)
    if not isinstance(skip, bool):
        raise ValueError(f"{source_id}: skip_tables must be true or false, not {skip!r}")
    exclude = options.get("exclude_sections", DEFAULT_EXCLUDE_SECTIONS)
    if not isinstance(exclude, (list, tuple)) or not all(isinstance(x, str) for x in exclude):
        raise ValueError(f"{source_id}: exclude_sections must be a list of heading names")
    return {"skip_tables": skip, "exclude_sections": tuple(exclude)}


def heading_key(text):
    """Heading text -> the form exclude_sections and REFERENCE_HEADINGS match.

    "F. References" and "References:" are both "references"; NCI's
    "Latest Updates to This Summary (05/12/2025)" loses its date.
    """
    key = clean(text).casefold()
    key = re.sub(r"\s*\([^()]*\d[^()]*\)$", "", key)
    key = re.sub(r"^(?:\d+(?:\.\d+)*\.?|[a-z][.)]|[ivxlc]+[.)])\s+", "", key)
    return key.strip(" :.")


# ============================================================ the units
_INVISIBLE = dict.fromkeys(map(ord, "­​‌‍﻿"), None)
# Link labels NCBI and NCI append to every reference - not part of the citation.
_LINK_LABEL_RE = re.compile(
    r"\[(?:PubMed(?: Abstract)?|PMC free article|Google Scholar|Cross ?Ref|DOI|"
    r"Free Full Text|Full Text)\b[^\]]*\]", re.IGNORECASE)
_ENUMERATOR_RE = re.compile(r"^\[?\(?\d{1,4}[.)\]]?$")


def text_of(text):
    """Collapse whitespace and drop soft hyphens / zero-width characters."""
    return clean(text.translate(_INVISIBLE))


def push_heading(headings, level, title):
    """Heading path after a heading of `level`: it replaces every heading
    at its level or deeper. Pure; headings is a tuple of (level, title)."""
    kept = tuple(h for h in headings if h[0] < level)
    return kept + ((level, title),)


class UnitCollector:
    """Receives headings and blocks in reading order, emits Units.

    Every loader drives one of these, so the rules that are the same for
    all formats live in one place: the section path, the locator, dropping
    excluded sections, turning everything under References into reference
    units, and keeping a text repeated at one place once (see _emit).
    """

    def __init__(self, title, exclude_sections=DEFAULT_EXCLUDE_SECTIONS, trust_levels=True):
        self.title = text_of(title or "")
        self.exclude = {heading_key(h) for h in exclude_sections}
        # False (PDFs, whose heading levels are guesses from font sizes): an
        # excluded or References heading governs only the text up to the
        # next heading of any level, so one ranked too high cannot swallow
        # - or turn into references - the rest of the document.
        self.trust_levels = trust_levels
        self.headings = ()          # ((level, title), ...), outermost first
        self.page = None            # set by the PDF loader, page by page
        self.units = []
        self._seen = set()
        self._label = ""            # a lone "1." (<dt>, <label>) waiting for its entry

    @property
    def section(self):
        return tuple(title for _, title in self.headings)

    @property
    def locator(self):
        if self.page is not None:
            return f"p. {self.page}"
        return self.headings[-1][1] if self.headings else self.title

    @property
    def _governing(self):
        """The headings whose name decides exclusion and reference-ness."""
        return self.headings if self.trust_levels else self.headings[-1:]

    @property
    def excluded(self):
        return any(heading_key(t) in self.exclude for _, t in self._governing)

    @property
    def in_references(self):
        return any(heading_key(t) in REFERENCE_HEADINGS for _, t in self._governing)

    def heading(self, level, title):
        title = text_of(title)
        if title:
            self.headings = push_heading(self.headings, level, title)
            self._label = ""

    def add(self, text, kind="prose"):
        """One paragraph / caption / reference. Prose under a References
        heading becomes a reference."""
        text = text_of(text)
        if not text or self.excluded:
            return
        if kind == "prose" and self.in_references:
            kind = "reference"
        if kind == "reference":
            if _ENUMERATOR_RE.match(text):
                self._label = text
                return
            text = text_of(_LINK_LABEL_RE.sub("", text))
            if self._label:
                text, self._label = f"{self._label} {text}", ""
        limit = config.CHUNK_TARGET_CHARS
        if kind != "prose" and len(text) > limit:
            # pack() never splits a caption or reference; a reference list
            # whose entries could not be told apart arrives as one block.
            for piece in split_long(text, limit):
                self._emit(piece, kind)
            return
        self._emit(text, kind)

    def add_table(self, rows, header=(), caption="", limit=None):
        """A table -> one table unit per row group (see table_texts)."""
        if self.excluded:
            return
        for text in table_texts(rows, header, caption, limit):
            self._emit(text, "table")

    def continue_last(self, text):
        """Append text to the last unit if it is prose that stopped
        mid-sentence - a paragraph running on across a PDF page break.
        Returns False (nothing changed) otherwise."""
        if not self.units or self.excluded:
            return False
        last = self.units[-1]
        if last.kind != "prose" or last.section != self.section or _ends_sentence(last.text):
            return False
        joined = join_lines([last.text, text])
        self.units[-1] = Unit(joined, last.section, last.locator, last.page, last.kind)
        return True

    def _emit(self, text, kind):
        """Keyed on the place the chunk ID hashes - the page for PDFs, the
        whole section path otherwise. Keyed on the innermost heading alone,
        the SCC "Cryosurgery" contraindications were dropped as repeats of
        the BCC ones. References keep the looser locator key: a citation in
        two sections' reference lists is one entry, not two."""
        if self.page is not None or kind == "reference":
            place = self.locator
        else:
            place = self.section
        key = (place, self.page, kind, text)
        if text and key not in self._seen:
            self._seen.add(key)
            self.units.append(Unit(text, self.section, self.locator, self.page, kind))


def _ends_sentence(text):
    return text.rstrip().endswith((".", "!", "?", ":", ";", '"', "”", ")", "]"))


# --------------------------------------------------------------- tables
def row_text(cells):
    """Cells -> "a | b | c". Trailing empty cells go; an all-empty row is ""."""
    cells = [text_of(c) for c in cells]
    while cells and not cells[-1]:
        cells.pop()
    return " | ".join(cells) if any(cells) else ""


def expand_spans(rows):
    """Rows of (text, rowspan, colspan) cells -> rows of cell texts, in
    which a cell spanning several rows is repeated in each of them.

    Without this, a stage cell spanning four rows ("III" over T1 N1, T2 N1,
    T3 N0, T3 N1) names only the first, and every later row shifts one
    column left, its T value under "Stage". A colspan pads with empty
    cells, so the next cell stays under its own header (a full-width
    footnote row stays one cell: row_text drops trailing empties).
    rowspan 0 means "to the last row given". Pure.
    """
    out, pending = [], {}      # column -> [text, rows still to fill]

    def carry(row, col, stop=None):
        """Fill row from col with the cells spanning down into it."""
        while col in pending or (stop is not None and col < stop):
            text, left = pending.pop(col, ("", 0))
            row.append(text)
            if left > 1:
                pending[col] = [text, left - 1]
            col += 1
        return col

    for r, cells in enumerate(rows):
        row, col = [], 0
        for text, rowspan, colspan in cells:
            col = carry(row, col)
            rowspan = len(rows) - r if rowspan == 0 else rowspan
            for k in range(colspan):
                row.append(text if k == 0 else "")
                if rowspan > 1:
                    pending[col] = [text if k == 0 else "", rowspan - 1]
                col += 1
        carry(row, col, stop=max(pending, default=-1) + 1)
        out.append(row)
    return out


def table_texts(rows, header=(), caption="", limit=None):
    """Rows -> texts of at most `limit` chars (config.CHUNK_TARGET_CHARS).

    Each text starts with the caption and the header row(s), so a row group
    on its own still says which table it is and what its columns mean. A
    single row longer than the limit is kept whole rather than cut mid-row.
    """
    limit = limit or config.CHUNK_TARGET_CHARS
    head = [text_of(caption)] if text_of(caption) else []
    head += [r for r in header if r]
    rows = [r for r in rows if r]
    if not rows:
        return ["\n".join(head)] if head else []

    groups, current = [], []
    size = sum(len(h) + 1 for h in head)
    for row in rows:
        if current and size + len(row) > limit:
            groups.append("\n".join(head + current))
            current, size = [], sum(len(h) + 1 for h in head)
        current.append(row)
        size += len(row) + 1
    groups.append("\n".join(head + current))
    return groups


# ====================================================== dispatch / load
def load(source):
    """sources.Source -> [chunking.Unit], by source.format."""
    if source.format == "pdf_textbook":
        raise ValueError(f"{source.id}: format pdf_textbook is read by ingest.build_chunks, "
                         "not by loaders.load")
    loader = LOADERS.get(source.format)
    if loader is None:
        raise ValueError(f"{source.id}: no loader for format {source.format!r} "
                         f"(have: {sorted(LOADERS)})")
    options = parse_options(source.options, source.id)
    return loader(Path(source.path), title=source.title, **options)


# ================================================================ HTML
_HEADING_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
_BLOCK_TAGS = _HEADING_TAGS | {
    "address", "article", "aside", "blockquote", "caption", "center", "dd", "details",
    "dialog", "div", "dl", "dt", "fieldset", "figcaption", "figure", "footer", "form",
    "header", "hgroup", "hr", "li", "main", "menu", "nav", "ol", "p", "pre", "section",
    "summary", "table", "ul"}
_PROSE_TAGS = {"p", "li", "dd", "dt", "blockquote", "pre", "address", "summary"}
# Not "form": ASP.NET pages - the Medicare Coverage Database's LCDs and
# billing articles - wrap the whole document in one <form>. Its controls
# go; the form itself is unwrapped unless it is a widget (strip_furniture).
_STRIP_TAGS = ["script", "style", "nav", "header", "footer", "aside", "noscript",
               "template", "button", "iframe", "svg", "select", "input", "textarea", "head"]
_STRIP_ROLES = {"navigation", "banner", "contentinfo", "search", "complementary"}
# A form with less text than this once its controls are gone, and no
# heading, is a widget (a search box, a newsletter sign-up), not content.
_WIDGET_FORM_CHARS = 200
# NCBI Bookshelf / PMC page furniture: sidebars (similar articles, recent
# activity), print/share widgets, "Go to:" links, the disclosure/copyright
# footer of a StatPearls chapter.
_STRIP_CLASSES = {"portlet", "bk_noprnt", "bk_prnt", "goto", "post-content", "pmc_page_banner",
                  "ncbi-alerts", "breadcrumb", "breadcrumbs", "share", "share-buttons",
                  "social-share", "print", "print-button", "skip-link", "back-to-top",
                  "similar-articles", "related-articles", "page-options", "fm-citation",
                  "pdq-hp-patient-toggle"}   # NCI "Go to Patient Version"
_STRIP_IDS = {"rightcolumn", "breadcrumb", "universal_header", "NCBIFooter_dynamic",
              "HTDisplay", "footer", "header", "print-header"}
_FURNITURE_HEADINGS = (
    "In This Section", "On This Page", "In this Page", "Similar articles",
    "Similar articles in PubMed", "Related information", "Related Items in Bookshelf",
    "Recent Activity", "Cite this Page", "Related links",
)
_MAIN_SELECTORS = ("main", "article", "#maincontent", ".document", "[role=main]")
_FIGURE_CLASSES = {"figure", "fig", "bk_fig", "iconblock"}
_TABLE_CLASSES = {"table", "table-wrap"}


def load_html(path, title="", skip_tables=False, exclude_sections=DEFAULT_EXCLUDE_SECTIONS):
    markup = Path(path).read_text(encoding="utf-8", errors="replace")
    return html_units(markup, title, skip_tables, exclude_sections)


def html_units(markup, title="", skip_tables=False, exclude_sections=DEFAULT_EXCLUDE_SECTIONS):
    """HTML string -> Units. BeautifulSoup with the stdlib "html.parser"."""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(markup, "html.parser")
    if not title and soup.title and soup.title.string:
        title = soup.title.string
    strip_furniture(soup)
    out = UnitCollector(title, tuple(exclude_sections) + _FURNITURE_HEADINGS)
    _HtmlWalker(out, skip_tables).container(main_container(soup))
    return out.units


def strip_furniture(soup):
    """Remove scripts, navigation, sidebars and widgets, in place.

    A <form> is removed only when it is a widget (is_widget_form); any
    other form is unwrapped, its content kept where it stood. Deleting
    every form left an MCD coverage page with no text at all."""
    for tag in soup.find_all(_STRIP_TAGS):
        tag.decompose()
    for form in soup.find_all("form"):
        if form.decomposed or form.attrs is None:
            continue
        if is_widget_form(form):
            form.decompose()
        else:
            form.unwrap()
    for tag in soup.find_all(True):
        if tag.decomposed or tag.attrs is None:
            continue
        classes = set(tag.get("class") or ())
        if (tag.get("role") in _STRIP_ROLES or tag.get("id") in _STRIP_IDS
                or classes & _STRIP_CLASSES or tag.has_attr("hidden")
                or tag.get("aria-hidden") == "true"):
            tag.decompose()


def is_widget_form(form):
    """A search box or a sign-up form: role="search", or (its controls
    already stripped) under _WIDGET_FORM_CHARS of text and no heading.
    The heading test keeps a short document that happens to sit in a form."""
    if form.get("role") == "search":
        return True
    return (len(form.get_text(strip=True)) < _WIDGET_FORM_CHARS
            and form.find(sorted(_HEADING_TAGS)) is None)


def main_container(soup):
    """The element that holds the document text: <main>, <article>,
    #maincontent or .document, else <body>, else the whole soup."""
    for selector in _MAIN_SELECTORS:
        found = soup.select_one(selector)
        if found is not None and found.get_text(strip=True):
            return found
    return soup.body or soup


def _classes(tag):
    return set(tag.get("class") or ())


def _has_body_text(el):
    """Does el hold any text besides its headings?"""
    in_headings = sum(len(h.get_text(strip=True)) for h in el.find_all(sorted(_HEADING_TAGS)))
    return len(el.get_text(strip=True)) > in_headings


def _is_footnote_mark(tag):
    """A superscript letter or symbol pointing at a table footnote ("Neck<sup>a</sup>").
    Kept, it glues onto the word ("Necka"). Numeric superscripts (cm2,
    citation numbers) stay."""
    if tag.name != "sup" and "sup" not in _classes(tag):
        return False
    return bool(re.fullmatch(r",?[a-z](?:\s*,\s*[a-z])*,?|[*†‡§¶#]+",
                             tag.get_text(strip=True)))


def _span(cell, attr):
    """A cell's rowspan / colspan as an int: 1 when missing or garbled,
    0 kept for rowspan ("to the end of the row group"), capped where the
    HTML spec caps them."""
    try:
        n = int(str(cell.get(attr, 1)).strip())
    except ValueError:
        return 1
    if attr == "rowspan":
        return min(n, 65534) if n >= 0 else 1
    return min(n, 1000) if n >= 1 else 1


def _is_caption(tag):
    return tag.name == "figcaption" or any(
        c == "caption" or c.endswith("-caption") or c.startswith("caption")
        for c in _classes(tag))


class _HtmlWalker:
    """Document-order walk. Containers (div, section, ul...) are recursed
    into; inline content between block children becomes a paragraph; each
    p / li / dd / blockquote is one paragraph of its own text, and blocks
    nested inside it (a figure inside a <p>, a sub-list inside an <li>) are
    handled after it, so no text is counted twice."""

    def __init__(self, out, skip_tables):
        self.out = out
        self.skip_tables = skip_tables

    def container(self, el):
        """HTML5 <section> / <article> scope their headings: when one closes,
        the heading path goes back to what it was before it opened. Without
        that, NCI's per-section <h6>References would nest under whichever
        <h3> came last (and vanish with it when that one is excluded).
        A section that held only headings (a title wrapper) keeps them."""
        scoped = el.name in ("section", "article") and _has_body_text(el)
        saved = self.out.headings
        run = []
        for child in el.children:
            if self._is_block(child):
                self._flush(run)
                run = []
                self.block(child)
            else:
                run.append(child)
        self._flush(run)
        if scoped:
            self.out.headings = saved

    def block(self, el):
        name, classes = el.name, _classes(el)
        if name in _HEADING_TAGS:
            self.out.heading(int(name[1]), self.text(el))
        elif (name == "table" or classes & _TABLE_CLASSES
              or name == "figure" and el.find("table") is not None):
            # HTML5 captions a table as <figure><table/><figcaption/></figure>;
            # figure() would keep the caption and lose every row.
            self.table(el)
        elif _is_caption(el):
            self.out.add(self.text(el), "caption")
        elif name == "figure" or classes & _FIGURE_CLASSES:
            self.figure(el)
        elif name in _PROSE_TAGS:
            nested = []
            text = self.inline(el, nested)
            if name == "li" and self.out.in_references:
                text = self._numbered(el, text)
            self.out.add(text)
            for child in nested:
                self.block(child)
        elif name != "hr":
            self.container(el)

    def figure(self, el):
        """Only the caption of a figure is text; the image, its alt text and
        "Enlarge" links are not. A Bookshelf figure block carries a "Figure"
        heading and a caption paragraph - joined into one caption."""
        captions = [c for c in el.find_all(True) if _is_caption(c)]
        ids = {id(c) for c in captions}
        captions = [c for c in captions if not any(id(p) in ids for p in c.parents)]
        if not captions:
            captions = el.find_all(["p"] + sorted(_HEADING_TAGS))
        text = " ".join(self.text(c) for c in captions)
        heading = el.find(sorted(_HEADING_TAGS))
        if heading is not None and self.text(heading) not in text:
            text = f"{self.text(heading)} {text}"
        self.out.add(text, "caption")

    def table(self, el):
        if self.skip_tables:
            return
        table = el if el.name == "table" else el.find("table")
        caption_parts = []
        if el is not table:   # wrapper (Bookshelf div.table, <figure>): its title and caption
            for part in el.find_all(sorted(_HEADING_TAGS) + ["div", "p", "figcaption"]):
                if (part.name in _HEADING_TAGS or _is_caption(part)) and not part.find_parent("table"):
                    caption_parts.append(self.text(part))
        if table is None:
            if any(_is_caption(c) for c in el.find_all(True)):
                self.out.add(" ".join(caption_parts), "caption")   # a table given as an image
            else:
                self.container(el)      # class="table" on a layout div, not a table
            return
        if table.caption:
            caption_parts.append(self.text(table.caption))
        rows = [tr for tr in table.find_all("tr") if tr.find_parent("table") is table]
        if rows and all(len(self._cells(tr)) <= 1 for tr in rows):
            self.container(table)      # a layout table, not data
            return
        # bs4 Tags compare equal by content, so membership is by identity.
        head = [tr for tr in rows if tr.find_parent("thead") is not None]
        foot = [tr for tr in rows if tr.find_parent("tfoot") is not None]
        taken = {id(tr) for tr in head + foot}
        body = [tr for tr in rows if id(tr) not in taken]
        if not head and body and all(c.name == "th" for c in self._cells(body[0])):
            head, body = body[:1], body[1:]
        grid = self._grid(rows)
        render = lambda trs: [row_text(grid[id(tr)]) for tr in trs]
        self.out.add_table(render(body) + render(foot), render(head), " ".join(caption_parts))

    def _grid(self, rows):
        """id(tr) -> its cell texts, rowspan and colspan expanded. A row
        group (thead, each tbody, tfoot) is expanded on its own: by the
        HTML spec a rowspan never reaches across one."""
        groups = {}
        for tr in rows:
            group = tr.find_parent(["thead", "tbody", "tfoot", "table"])
            groups.setdefault(id(group), []).append(tr)
        grid = {}
        for trs in groups.values():
            cells = [[(self.text(c), _span(c, "rowspan"), _span(c, "colspan"))
                      for c in self._cells(tr)] for tr in trs]
            for tr, texts in zip(trs, expand_spans(cells)):
                grid[id(tr)] = texts
        return grid

    @staticmethod
    def _cells(tr):
        return tr.find_all(["th", "td"], recursive=False)

    def _numbered(self, li, text):
        """An entry of a numbered reference list keeps its number, so the
        "[18]" in the prose can be looked up."""
        parent = li.parent
        if parent is None or parent.name != "ol" or re.match(r"^\[?\d", text):
            return text
        try:
            start = int(parent.get("start", 1))
        except ValueError:
            start = 1
        siblings = parent.find_all("li", recursive=False)
        index = next((i for i, s in enumerate(siblings) if s is li), None)
        return text if index is None else f"{start + index}. {text}"

    # ---- text
    def _is_block(self, node):
        name = getattr(node, "name", None)
        return name is not None and (name in _BLOCK_TAGS or _is_caption(node))

    def _flush(self, run):
        nested = []
        text = "".join(self.inline(node, nested) for node in run)
        self.out.add(text)
        for child in nested:
            self.block(child)

    def inline(self, node, nested):
        """Text of node, without nested block elements (those go to
        `nested`, to be handled after the paragraph they sit in)."""
        from bs4.element import NavigableString, PreformattedString
        if isinstance(node, NavigableString):
            return "" if isinstance(node, PreformattedString) else str(node)
        if node.name == "br":
            return " "
        if _is_footnote_mark(node):
            return ""
        parts = []
        for child in node.children:
            if self._is_block(child):
                nested.append(child)
            else:
                parts.append(self.inline(child, nested))
        return "".join(parts)

    @classmethod
    def text(cls, el):
        """All text of el; <br> and block boundaries become spaces. A figure
        inside (NCI puts staging drawings in table cells) gives only its
        caption, not its "Enlarge" link."""
        return text_of(cls._text(el))

    @classmethod
    def _text(cls, node):
        from bs4.element import NavigableString, PreformattedString
        if isinstance(node, NavigableString):
            return "" if isinstance(node, PreformattedString) else str(node)
        if node.name == "br":
            return " "
        if _is_footnote_mark(node):
            return ""
        if node.name == "figure" or _classes(node) & _FIGURE_CLASSES:
            captions = [c for c in node.find_all(True) if _is_caption(c)][:1]
            return " " + " ".join(cls._text(c) for c in captions) + " "
        inner = "".join(cls._text(child) for child in node.children)
        if node.name in _BLOCK_TAGS or node.name in ("td", "th"):
            return f" {inner} "
        return inner


# ============================================================ JATS nxml
_XML_ENTITIES = {"amp", "lt", "gt", "quot", "apos"}
_JATS_SECTIONS = {"sec", "app", "ack", "ref-list", "glossary", "notes", "fn-group",
                  "boxed-text", "bio"}
_JATS_DEFAULT_TITLES = {"ref-list": "References", "ack": "Acknowledgments",
                        "glossary": "Glossary"}
_JATS_BLOCKS = {"fig", "fig-group", "table-wrap", "table-wrap-group", "list", "disp-quote",
                "boxed-text", "def-list", "sec", "statement", "speech", "verse-group",
                "code", "preformat", "graphic", "media", "disp-formula-group"}
_JATS_SKIP = {"title", "label", "permissions", "graphic", "media", "supplementary-material",
              "object-id", "sec-meta", "kwd-group", "alt-text", "long-desc", "tex-math",
              "annotation", "inline-graphic", "x"}
_SKIP_ABSTRACTS = {"toc", "graphical", "teaser", "web-summary"}


def load_nxml(path, title="", skip_tables=False, exclude_sections=DEFAULT_EXCLUDE_SECTIONS):
    xml_text = Path(path).read_text(encoding="utf-8", errors="replace")
    return nxml_units(xml_text, title, skip_tables, exclude_sections)


def _resolve_entities(xml_text):
    """Named HTML entities (&nbsp; &mdash;) -> characters. JATS defines them
    in its DTD, which ElementTree never reads, so they would fail to parse."""
    def replace(m):
        name = m.group(1)
        if name in _XML_ENTITIES or name not in name2codepoint:
            return m.group(0)
        return f"&#{name2codepoint[name]};"
    return re.sub(r"&([A-Za-z][A-Za-z0-9]*);", replace, xml_text)


def _local(tag):
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def nxml_units(xml_text, title="", skip_tables=False, exclude_sections=DEFAULT_EXCLUDE_SECTIONS):
    """JATS / BITS XML string (PMC articles, NCBI LitArch books) -> Units.

    The article title is the root of the section path (as an HTML <h1> is),
    then "Abstract", then the nested <sec>/<title>s of the body and back
    matter. Namespaced tags are reduced to their local names.
    """
    root = ET.fromstring(_resolve_entities(xml_text))
    for el in root.iter():
        el.tag = _local(el.tag)
    doc_title = _jats_title(root) or title
    out = UnitCollector(doc_title or title, exclude_sections)
    walker = _JatsWalker(out, skip_tables)
    out.heading(0, doc_title)
    base = out.headings

    for abstract in root.iter("abstract"):
        if abstract.get("abstract-type") in _SKIP_ABSTRACTS:
            continue
        out.headings = base
        out.heading(1, "Abstract")
        walker.walk(abstract, 2)
    for part in ("body", "back"):
        for el in root.iter(part):
            out.headings = base
            walker.walk(el, 1)
    return out.units


def _jats_title(root):
    for path in (".//book-part-meta/title-group/title", ".//article-meta/title-group/article-title",
                 ".//article-title", ".//title-group/title", ".//book-title"):
        el = root.find(path)
        if el is not None and _jats_text(el):
            return _jats_text(el)
    return ""


def _jats_text(el, nested=None):
    """Text of an element. With `nested` given, nested blocks (figures,
    tables, lists) are left out and collected there instead."""
    if el is None:
        return ""
    parts = [el.text or ""]
    for child in el:
        if child.tag in _JATS_SKIP and child.tag != "title":
            pass
        elif nested is not None and child.tag in _JATS_BLOCKS:
            nested.append(child)
        else:
            inner = _jats_text(child, nested)
            parts.append(f" {inner} " if child.tag in _NAME_PARTS else inner)
        parts.append(child.tail or "")
    return text_of("".join(parts))


def _fix_punctuation_spacing(text):
    return text_of(re.sub(r"\s+([,.;:)\]])", r"\1", re.sub(r"([(\[])\s+", r"\1", text)))


class _JatsWalker:
    def __init__(self, out, skip_tables):
        self.out = out
        self.skip_tables = skip_tables

    def walk(self, el, depth):
        for child in el:
            tag = child.tag
            if tag in _JATS_SKIP or tag in ("abstract", "trans-abstract"):
                continue
            if tag in _JATS_SECTIONS:
                self.section(child, depth)
            elif tag == "p":
                nested = []
                self.out.add(_jats_text(child, nested))
                for block in nested:
                    self.walk_block(block, depth)
            elif tag == "ref":
                self.out.add(citation_text(child), "reference")
            elif tag in ("fig", "table-wrap"):
                self.walk_block(child, depth)
            elif tag in ("list-item", "def-item", "def", "term") and not any(
                    c.tag in ("p", "list") for c in child):
                self.out.add(_jats_text(child))
            else:
                self.walk(child, depth)

    def walk_block(self, el, depth):
        if el.tag == "fig":
            self.out.add(" ".join(filter(None, (_jats_text(el.find("label")),
                                                _jats_text(el.find("caption"))))), "caption")
        elif el.tag == "table-wrap":
            self.table(el)
        else:
            wrapper = ET.Element("wrapper")
            wrapper.append(el)
            self.walk(wrapper, depth)

    def section(self, el, depth):
        saved = self.out.headings
        heading = _jats_text(el.find("title")) or _JATS_DEFAULT_TITLES.get(el.tag, "")
        if heading:
            self.out.heading(depth, heading)
        self.walk(el, depth + 1 if heading else depth)
        self.out.headings = saved

    def table(self, wrap):
        if self.skip_tables:
            return
        caption = " ".join(filter(None, (_jats_text(wrap.find("label")),
                                         _jats_text(wrap.find("caption")))))
        table = wrap.find(".//table")
        if table is None:      # the table is only an image
            self.out.add(caption, "caption")
            return
        head, body = [], []
        for group in table.iter():
            if group.tag != "tr":
                continue
            cells = [_jats_text(c) for c in group if c.tag in ("th", "td")]
            (head if _in_thead(table, group) else body).append(row_text(cells))
        foot = [_jats_text(fn) for fn in wrap.iter("table-wrap-foot")]
        self.out.add_table(body + foot, head, caption)


def _in_thead(table, tr):
    return any(tr in list(thead.iter("tr")) for thead in table.iter("thead"))


def citation_text(ref):
    """<ref> -> "12. Mohs FE. Chemosurgery. Arch Surg. 1941;42:279-95."

    mixed-citation already holds its punctuation; element-citation is bare
    fields, so it is assembled in Vancouver order.
    """
    label = _jats_text(ref.find("label")).rstrip(".")
    cite = next((c for c in ref if c.tag in ("mixed-citation", "element-citation",
                                             "citation", "nlm-citation")), None)
    if cite is None:
        text = _jats_text(ref)
    elif cite.tag == "element-citation":
        text = _element_citation(cite)
    else:
        text = _fix_punctuation_spacing(_jats_text_without(cite, {"pub-id", "comment"}))
    text = text_of(text)
    if label and not text.startswith(label):
        text = f"{label}. {text}"
    return text


_NAME_PARTS = ("name", "surname", "given-names", "collab", "string-name", "suffix", "prefix")


def _jats_text_without(el, skip):
    """Text of a mixed-citation minus `skip` (pub-ids). Its own text nodes
    hold the punctuation; only name parts need a space between them
    (<surname>Mohs</surname><given-names>FE</given-names> -> "Mohs FE")."""
    parts = [el.text or ""]
    for child in el:
        if child.tag not in skip:
            inner = _jats_text_without(child, skip)
            parts.append(f" {inner} " if child.tag in _NAME_PARTS else inner)
        parts.append(child.tail or "")
    return "".join(parts)


def _element_citation(cite):
    names = []
    for group in cite.iter("person-group"):
        for name in group:
            if name.tag == "name":
                names.append(text_of(f"{_jats_text(name.find('surname'))} "
                                     f"{_jats_text(name.find('given-names'))}"))
            elif name.tag == "collab":
                names.append(_jats_text(name))
            elif name.tag == "etal":
                names.append("et al")
    get = lambda tag: _jats_text(cite.find(tag))
    pages = "-".join(filter(None, (get("fpage"), get("lpage"))))
    where = get("year")
    if get("volume"):
        where += f";{get('volume')}" + (f"({get('issue')})" if get("issue") else "")
    if pages:
        where += f":{pages}"
    pieces = [", ".join(names), get("article-title") or get("chapter-title"),
              get("source"), where]
    return " ".join(f"{p.rstrip('.')}." for p in pieces if p)


# ================================================================ DOCX
def load_docx(path, title="", skip_tables=False, exclude_sections=DEFAULT_EXCLUDE_SECTIONS):
    """Word document -> Units. Title / Heading N styles drive the section
    path; paragraphs and tables are read in body order (python-docx's
    doc.paragraphs and doc.tables would separate them), content controls
    and custom-XML wrappers included (docx_blocks). Text comes from
    docx_text, not python-docx's Paragraph.text / _Cell.text, which miss
    tracked insertions, content controls and nested tables."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = docx.Document(str(path))
    out = UnitCollector(title, exclude_sections)
    for block in docx_blocks(doc.element.body):
        if _local(block.tag) == "p":
            paragraph = Paragraph(block, doc)
            style = paragraph.style.name if paragraph.style is not None else ""
            level = docx_heading_level(style)
            text = docx_text(block)
            if level is not None:
                out.heading(level, text)
            elif style.lower() == "caption":
                out.add(text, "caption")
            else:
                out.add(text)
        elif not skip_tables:
            rows = [row_text(_docx_cells(row)) for row in Table(block, doc).rows]
            out.add_table(rows[1:], rows[:1])
    return out.units


_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC_NS = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
# Never visible text: deleted / moved-away revisions, the fallback copy of
# an mc:AlternateContent (the same text box again, as VML), and property
# elements (pPr's w:tabs holds w:tab tab stops, not tabs).
_DOCX_SKIP = {_W_NS + t for t in ("del", "moveFrom", "pPr", "rPr", "sdtPr", "sdtEndPr",
                                  "tblPr", "trPr", "tcPr")} | {_MC_NS + "Fallback"}
_DOCX_CHARS = {_W_NS + "tab": " ", _W_NS + "ptab": " ", _W_NS + "br": " ", _W_NS + "cr": " ",
               _W_NS + "noBreakHyphen": "-"}


def docx_blocks(parent):
    """The w:p and w:tbl elements of a body or a table cell, in order,
    looking through content controls (w:sdt -> w:sdtContent) and
    w:customXml, which can wrap whole headings, paragraphs and tables."""
    for child in parent.iterchildren():
        tag = child.tag
        if tag in (_W_NS + "p", _W_NS + "tbl"):
            yield child
        elif tag == _W_NS + "sdt" and not _showing_placeholder(child):
            for content in child.iterchildren(_W_NS + "sdtContent"):
                yield from docx_blocks(content)
        elif tag == _W_NS + "customXml":
            yield from docx_blocks(child)


def _showing_placeholder(sdt):
    """An empty content control shows its prompt ("Click or tap here to
    enter text.") as real runs - not document text."""
    return sdt.find(f"{_W_NS}sdtPr/{_W_NS}showingPlcHdr") is not None


def docx_text(el):
    """Text of a paragraph as Word shows it with changes accepted: every
    w:t below it - inside runs, hyperlinks, tracked insertions, inline
    content controls, smart tags, fields - minus _DOCX_SKIP. So the tracked
    replacement "Max dose is <del>5</del><ins>7</ins> mg/kg" reads
    "Max dose is 7 mg/kg", not python-docx's "Max dose is  mg/kg"."""
    parts = []
    for child in el.iterchildren():
        tag = child.tag
        if not isinstance(tag, str) or tag in _DOCX_SKIP:
            continue
        if tag == _W_NS + "sdt" and _showing_placeholder(child):
            continue
        if tag == _W_NS + "t":
            parts.append(child.text or "")
        elif tag in _DOCX_CHARS:
            parts.append(_DOCX_CHARS[tag])
        else:
            parts.append(docx_text(child))
    return "".join(parts)


def docx_heading_level(style_name):
    """"Title" -> 0, "Heading 2" -> 2, anything else -> None."""
    if style_name == "Title":
        return 0
    m = re.match(r"^Heading (\d)$", style_name or "")
    return int(m.group(1)) if m else None


def _docx_cells(row):
    """Cell texts of a row; a merged cell is listed once, not per column."""
    cells, previous = [], None
    for cell in row.cells:
        if cell._tc is not previous:
            cells.append(docx_cell_text(cell._tc))
        previous = cell._tc
    return cells


def docx_cell_text(tc):
    """A w:tc's text: its paragraphs, and a table nested in it as
    "a, b; c, d" - rows joined by "; ", cells by ", ". Not " | ": that
    would read as more columns of the outer table."""
    parts = []
    for block in docx_blocks(tc):
        if block.tag == _W_NS + "p":
            parts.append(docx_text(block))
        else:
            rows = []
            for tr in block.iterchildren(_W_NS + "tr"):
                cells = [text_of(docx_cell_text(c)) for c in tr.iterchildren(_W_NS + "tc")]
                rows.append(", ".join(c for c in cells if c))
            parts.append("; ".join(r for r in rows if r))
    return " ".join(p for p in parts if p.strip())


# ============================================================ Markdown
_ATX_RE = re.compile(r"^ {0,3}(#{1,6})\s+(.*?)(?:\s+#+)?\s*$")
_FENCE_RE = re.compile(r"^ {0,3}(`{3,}|~{3,})")
_LIST_ITEM_RE = re.compile(r"^\s*(?:[-*+]|\d{1,9}[.)])\s+")
_THEMATIC_RE = re.compile(r"^ {0,3}(?:(?:-\s*){3,}|(?:\*\s*){3,}|(?:_\s*){3,})$")
_SETEXT_RE = re.compile(r"^ {0,3}(=+|-+)\s*$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{1,}:?\s*(?:\|\s*:?-{1,}:?\s*)*\|?\s*$")


def load_markdown(path, title="", skip_tables=False, exclude_sections=DEFAULT_EXCLUDE_SECTIONS):
    text = Path(path).read_text(encoding="utf-8", errors="replace")
    return markdown_units(text, title, skip_tables, exclude_sections)


def md_inline(text):
    """Markdown inline syntax -> plain text: images dropped, links keep
    their text, **bold** / __bold__ / `code` lose their markers, HTML tags go."""
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\[[^\]]*\]", r"\1", text)
    text = re.sub(r"<!--.*?-->", "", text)
    text = re.sub(r"</?[A-Za-z][^>]*>", " ", text)
    text = re.sub(r"(\*\*|__)(.+?)\1", r"\2", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text


def md_table_row(line):
    """"| a | b \\| c |" -> ["a", "b | c"]."""
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    return [md_inline(c.replace("\\|", "|")) for c in re.split(r"(?<!\\)\|", line)]


# Front-matter lines when PyYAML is not installed: "key: value", "- item",
# indented continuations, blanks.
_YAML_LINE_RE = re.compile(r"^(?:\s*[A-Za-z_][\w-]*\s*:(?:\s|$)|\s*-\s|\s+\S|\s*$)")


def _skip_front_matter(lines):
    """Lines minus a leading YAML front-matter block. A "---" first line
    alone proves nothing - a plain-text protocol may open with a rule - so
    the block goes only if it is a YAML mapping (title:, author: ...)."""
    if lines and lines[0].strip() == "---":
        for i, line in enumerate(lines[1:60], start=1):
            if line.strip() in ("---", "..."):
                return lines[i + 1:] if is_yaml_mapping(lines[1:i]) else lines
    return lines


def is_yaml_mapping(lines):
    """Do these lines parse as a YAML mapping? PyYAML is imported here,
    not at the top: it is an ingest-only dependency (requirements-dev.txt),
    never the app's. Without it, every non-blank line must look like YAML
    (_YAML_LINE_RE)."""
    try:
        import yaml
    except ImportError:
        return any(l.strip() for l in lines) and all(_YAML_LINE_RE.match(l) for l in lines)
    try:
        return isinstance(yaml.safe_load("\n".join(lines)), dict)
    except yaml.YAMLError:
        return False


def markdown_units(text, title="", skip_tables=False, exclude_sections=DEFAULT_EXCLUDE_SECTIONS):
    """Markdown (or plain text) -> Units.

    ATX (#) and setext (===) headings set the section path. Paragraphs are
    separated by blank lines; consecutive list items stay one paragraph (a
    list is one thought), except under References, where each item is one
    entry. Pipe tables become table units; fenced code is kept as text.
    A leading byte-order mark (Notepad's "UTF-8 with BOM") is dropped, or
    the first line would be neither "---" nor a "#" heading.
    """
    out = UnitCollector(title, exclude_sections)
    lines = _skip_front_matter(text.lstrip("\ufeff").splitlines())
    para = []

    def flush():
        if para:
            text = " ".join(line.strip() for line in para)
            if out.in_references:       # "- Smith J. ..." -> "Smith J. ..."
                text = re.sub(r"^[-*+]\s+", "", text)
            out.add(md_inline(text))
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        fence = _FENCE_RE.match(line)
        if fence:
            flush()
            marker, code = fence.group(1)[0] * 3, []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith(marker):
                code.append(lines[i])
                i += 1
            out.add(" ".join(code))
            i += 1
            continue
        heading = _ATX_RE.match(line)
        if heading:
            flush()
            out.heading(len(heading.group(1)), md_inline(heading.group(2)))
        elif not line.strip():
            flush()
        elif "|" in line and i + 1 < len(lines) and _TABLE_SEP_RE.match(lines[i + 1]) \
                and "|" in lines[i + 1] and "-" in lines[i + 1]:
            flush()
            header = [row_text(md_table_row(line))]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip() and "|" in lines[i]:
                rows.append(row_text(md_table_row(lines[i])))
                i += 1
            if not skip_tables:
                out.add_table(rows, header)
            continue
        elif _SETEXT_RE.match(line) and len(para) == 1 and not _LIST_ITEM_RE.match(para[0]):
            out.heading(1 if line.strip()[0] == "=" else 2, md_inline(para.pop().strip()))
        elif _THEMATIC_RE.match(line):
            flush()
        else:
            if _LIST_ITEM_RE.match(line) and out.in_references:
                flush()
            para.append(line)
        i += 1
    flush()
    return out.units


# ================================================================= PDF
# Running headers and footers live in the outer band of the page. CMS
# manuals put theirs about 12% up from the bottom edge (the "Revision Date"
# footer of the NCCI manual), so the band is wider than the usual 7%.
# A band that wide also holds the first and last few body lines, table rows
# included, so a line there is only a furniture candidate if it is also
# alone on its baseline (a table cell never is) and outside the page's
# main text block: among the first / last EDGE_ROWS rows and set off from
# the body by a gap of more than EDGE_GAP line heights. A candidate is
# furniture if it is a page number, or if its text, digits removed, repeats
# on at least 30% of pages - body text never does. Texts that short once
# the digits are gone ("2 g" -> "g", "T1" -> "t") are left out of the
# repetition count: different cells would collapse into one "header".
PDF_MARGIN = 0.15
PDF_REPEAT_FRACTION = 0.30
EDGE_ROWS = 3                 # a running header / footer is at most this many rows
EDGE_GAP = 1.0                # gap, in body line heights, that sets one off
MIN_FURNITURE_CHARS = 4       # shorter normalized texts are never "repeated furniture"
HEADING_SIZE_RATIO = 1.15     # this much larger than body text = a heading
MAX_HEADING_CHARS = 150
PARAGRAPH_GAP = 0.5           # vertical gap, in line heights, that ends a paragraph
# Valid roman numerals up to 399 only: no front matter runs longer, and
# leaving out m and d keeps "mL", "md", "mild", "lid", "IM", "ID" out.
_ROMAN = r"(?=[ivxlc])c{0,3}(?:xc|xl|l?x{0,3})(?:ix|iv|v?i{0,3})"
_PAGE_NUMBER_RE = re.compile(rf"^(?:page\s+)?(?:\d{{1,4}}|{_ROMAN})(?:\s*(?:of|/)\s*\d{{1,4}})?$",
                             re.IGNORECASE)
_DOT_LEADER_RE = re.compile(r"(?:\.\s?){6,}|(?:…\s?){3,}")
_ENTRY_START_RE = re.compile(r"^\s*(?:\[\d{1,4}\]|\d{1,3}\.)\s+\S")
# Line-end hyphens kept when the document gives no hint either way: after
# these, a hyphen is almost always part of a compound ("ill-defined").
_COMPOUND_PREFIXES = {"ill", "non", "self", "well", "semi", "intra", "post", "cross", "long", "short"}


def load_pdf(path, title="", skip_tables=False, exclude_sections=DEFAULT_EXCLUDE_SECTIONS):
    """A generic PDF (a policy manual, a guideline) -> Units, one or more per
    paragraph, with page = the 1-based page it starts on."""
    import pymupdf
    flags = pymupdf.TEXT_PRESERVE_WHITESPACE | pymupdf.TEXT_MEDIABOX_CLIP
    with pymupdf.open(str(path)) as doc:
        pages = [(page.rect.height, pdf_lines(page.get_text("dict", flags=flags)))
                 for page in doc]
    return pdf_units(pages, title, exclude_sections)


def _is_bold(span):
    return bool(span.get("flags", 0) & 16) or "bold" in span.get("font", "").lower()


def pdf_lines(page_dict):
    """page.get_text("dict") -> [line dict] in content-stream order.

    A line is {text, size, bold, y0, y1}; size is the size most of its
    characters have, bold means every visible span is bold. Whitespace-only
    lines are kept with text "" - they separate paragraphs. Rotated text
    (margin stamps) is skipped.
    """
    lines = []
    for block in page_dict.get("blocks", []):
        if block.get("type", 0) != 0:
            continue
        for line in block.get("lines", []):
            if abs(line.get("dir", (1, 0))[1]) > 0.1:
                continue
            spans = line.get("spans", [])
            visible = [s for s in spans if s["text"].strip()]
            text = "".join(s["text"] for s in spans)
            sizes = Counter()
            for s in visible:
                sizes[round(s["size"], 1)] += len(s["text"].strip())
            x0, y0, x1, y1 = line["bbox"]
            lines.append({
                "text": text if visible else "",
                "size": sizes.most_common(1)[0][0] if sizes else 0.0,
                "bold": bool(visible) and all(_is_bold(s) for s in visible),
                "y0": y0, "y1": y1,
            })
    return lines


def _norm_furniture(text):
    return " ".join(re.sub(r"\d+", "", text).lower().split())


def _in_margin(line, height):
    middle = (line["y0"] + line["y1"]) / 2
    return middle < height * PDF_MARGIN or middle > height * (1 - PDF_MARGIN)


def baseline_rows(lines):
    """Indices of the non-empty lines, grouped by baseline (y0 within
    2 pt - the same test pdf_paragraphs uses), top to bottom."""
    rows = []
    for i in sorted((i for i, l in enumerate(lines) if l["text"].strip()),
                    key=lambda i: lines[i]["y0"]):
        if rows and abs(lines[i]["y0"] - lines[rows[-1][0]]["y0"]) < 2:
            rows[-1].append(i)
        else:
            rows.append([i])
    return rows


def _edge_rows(lines, rows):
    """Row numbers outside the main text block: the first (last) k <=
    EDGE_ROWS rows when a gap of more than EDGE_GAP body line heights
    follows (precedes) them - the largest such k. A page of no more than
    EDGE_ROWS rows has no body block, so all of it counts."""
    def top(row):
        return min(lines[i]["y0"] for i in row)

    def bottom(row):
        return max(lines[i]["y1"] for i in row)

    def set_off(upper, lower, body):
        height = max(max(lines[i]["y1"] - lines[i]["y0"] for i in body), 1.0)
        return top(lower) - bottom(upper) > EDGE_GAP * height

    n, edge = len(rows), set()
    head = [k for k in range(1, min(EDGE_ROWS, n) + 1)
            if k == n or set_off(rows[k - 1], rows[k], rows[k])]
    foot = [k for k in range(1, min(EDGE_ROWS, n) + 1)
            if k == n or set_off(rows[n - k - 1], rows[n - k], rows[n - k - 1])]
    if head:
        edge.update(range(max(head)))
    if foot:
        edge.update(range(n - max(foot), n))
    return edge


def furniture_candidates(lines, height):
    """Indices of the lines that may be a running header, footer or page
    number: in the margin band, alone on their baseline, outside the main
    text block (see PDF_MARGIN). Pure."""
    rows = baseline_rows(lines)
    return {rows[r][0] for r in _edge_rows(lines, rows)
            if len(rows[r]) == 1 and _in_margin(lines[rows[r][0]], height)}


def furniture_texts(pages):
    """Normalized texts of running headers / footers: candidate lines
    (furniture_candidates) whose text, digits removed, is at least
    MIN_FURNITURE_CHARS long and repeats on >= 30% of pages (and at least
    2). pages = [(page_height, lines)]."""
    counts = Counter()
    for height, lines in pages:
        counts.update({_norm_furniture(lines[i]["text"])
                       for i in furniture_candidates(lines, height)})
    needed = max(2, ceil(PDF_REPEAT_FRACTION * len(pages)))
    return {text for text, n in counts.items()
            if len(text) >= MIN_FURNITURE_CHARS and n >= needed}


def _is_furniture(text, candidate, furniture):
    """A table-of-contents line anywhere; a candidate line (see
    furniture_candidates) that is a page number or a repeated header."""
    text = text.strip()
    if _DOT_LEADER_RE.search(text):
        return True
    if not candidate:
        return False
    return _norm_furniture(text) in furniture or bool(_PAGE_NUMBER_RE.match(text))


def body_size(pages):
    """The most common font size, weighted by characters: the body text."""
    sizes = Counter()
    for _, lines in pages:
        for line in lines:
            if line["text"].strip():
                sizes[line["size"]] += len(line["text"].strip())
    return sizes.most_common(1)[0][0] if sizes else 0.0


def pdf_paragraphs(lines, body, known_headings=frozenset()):
    """Lines of one page -> paragraphs (lists of lines).

    A paragraph ends at a blank line, at a vertical gap of more than half a
    line, when the text jumps back up the page (a new column), when the
    font size class changes (heading vs body), and before a line that
    starts a numbered entry ("[3] ...", "12. ..."). Lines on one baseline
    (a bullet and its text) belong together.

    A bold line that is nothing but a known section name ("References",
    an excluded heading) is a paragraph of its own even with no gap below
    it - books often set "References" straight above the first entry.
    """
    def known(line):
        return ((line["bold"] or _is_large(line, body))
                and heading_key(line["text"]) in known_headings)

    paragraphs, current = [], []
    for line in lines:
        if not line["text"].strip():
            if current:
                paragraphs.append(current)
            current = []
            continue
        if current:
            prev = current[-1]
            same_baseline = abs(line["y0"] - prev["y0"]) < 2
            height = max(prev["y1"] - prev["y0"], 1.0)
            new = not same_baseline and (
                line["y0"] - prev["y1"] > PARAGRAPH_GAP * height
                or line["y0"] < prev["y0"] - 2
                or _is_large(line, body) != _is_large(prev, body)
                or _ENTRY_START_RE.match(line["text"])
                or known(line) or known(prev))
            if new:
                paragraphs.append(current)
                current = []
        current.append(line)
    if current:
        paragraphs.append(current)
    return paragraphs


def _is_large(line, body):
    return body > 0 and line["size"] >= HEADING_SIZE_RATIO * body


def vocabulary(pages):
    """Lowercased words of the document, hyphenated compounds included -
    the evidence join_lines uses to decide about a line-end hyphen."""
    words = set()
    for _, lines in pages:
        for line in lines:
            words.update(w.lower() for w in re.findall(r"[A-Za-z]+(?:-[A-Za-z]+)*", line["text"]))
    return words


_SUSPENDED = {"and", "or", "to", "nor", "and/or"}


def join_lines(texts, words=frozenset()):
    """Lines -> one string, mending words broken at a line end.

      "recon-" + "struction..."  -> "reconstruction" (a hyphenated break)
      "ill-" + "defined"         -> "ill-defined"   (the document writes it so,
                                    or a compound prefix like ill-, non-)
      "99151-" + "99153"         -> "99151-99153"   (a range: hyphen, no space)
      "pre-" + "and post-"       -> "pre- and post-" (a suspended hyphen)
    """
    out = ""
    for text in texts:
        text = text.strip()
        if not text:
            continue
        hyphen = re.search(r"([A-Za-z0-9]+)-$", out)
        follow = re.match(r"([A-Za-z0-9/]+)", text)
        if not (hyphen and follow) or follow.group(1).lower() in _SUSPENDED:
            out = f"{out} {text}" if out else text
            continue
        head, tail = hyphen.group(1).lower(), follow.group(1)
        if not tail.islower() or not head.isalpha():
            out += text          # "99151-99153", "Procedure-to-Procedure"
            continue
        keep = (f"{head}-{tail}" in words and f"{head}{tail}" not in words) or (
            f"{head}{tail}" not in words and head in _COMPOUND_PREFIXES)
        out = out + text if keep else out[:-1] + text
    return text_of(out)


def _paragraph_kind(text, lines, body):
    """"heading", "caption" or "prose" for one PDF paragraph."""
    if CAPTION_RE.match(text) and re.match(r"^\S+\.?\s*\d", text):
        return "caption"
    if len(text) > MAX_HEADING_CHARS or not re.search(r"[A-Za-z]", text):
        return "prose"
    if all(_is_large(l, body) for l in lines):
        return "heading"
    if all(l["bold"] for l in lines) and not text.endswith((".", ",", ";", ":")):
        return "heading"
    return "prose"


def pdf_units(pages, title="", exclude_sections=DEFAULT_EXCLUDE_SECTIONS):
    """[(page_height, lines)] -> Units. Pure, so tests can feed it lines.

    Heading levels come from size rank: the largest heading size is level
    1, the next 2, ...; bold headings at body size come last.
    """
    body = body_size(pages)
    furniture = furniture_texts(pages)
    words = vocabulary(pages)
    known = REFERENCE_HEADINGS | {heading_key(h) for h in exclude_sections}

    parsed = []        # per page: [(kind, text, size)]
    large = []         # (size, page) of every large-font heading
    for page_number, (height, lines) in enumerate(pages, start=1):
        candidates = furniture_candidates(lines, height)
        kept = [l for i, l in enumerate(lines) if not (
            l["text"].strip() and _is_furniture(l["text"], i in candidates, furniture))]
        items = []
        for para in pdf_paragraphs(kept, body, known):
            text = join_lines([l["text"] for l in para], words)
            kind = _paragraph_kind(text, para, body)
            size = max(l["size"] for l in para)
            if kind == "heading" and _is_large(para[0], body):
                large.append((size, page_number))
            items.append((kind, text, size))
        parsed.append(items)

    levels = heading_levels(large)
    body_level = max(levels.values(), default=0) + 1
    out = UnitCollector(title, exclude_sections, trust_levels=False)
    for page_number, items in enumerate(parsed, start=1):
        out.page = page_number
        for i, (kind, text, size) in enumerate(items):
            if kind == "heading":
                out.heading(levels.get(size, body_level), text)
            elif i == 0 and kind == "prose" and _continues(text) and out.continue_last(text):
                pass     # a paragraph running on from the previous page stays whole
            else:
                out.add(text, kind)
    return out.units


def heading_levels(headings):
    """[(size, page)] of the large-font headings -> {size: level}.

    Bigger is higher: the largest size is level 1. Except a size used only
    once after page 1 - a one-off banner like a book's "Contributors" page -
    which ranks with the next size down; ranked by size alone it would sit
    above every chapter for the rest of the document. Once on page 1 is the
    document title and keeps its place.
    """
    counts = Counter(size for size, _ in headings)
    first_page = {}
    for size, page in headings:
        first_page.setdefault(size, page)
    ranked = sorted((s for s in counts if counts[s] > 1 or first_page[s] == 1), reverse=True)
    levels = {s: i + 1 for i, s in enumerate(ranked)}
    for s in counts:
        if s not in levels:
            below = [r for r in ranked if r < s]
            levels[s] = levels[below[0]] if below else len(ranked) + 1
    return levels


def _continues(text):
    """Could this first paragraph of a page be the rest of the previous
    one? Not if it starts a numbered entry or a bullet."""
    return not _ENTRY_START_RE.match(text) and not text.startswith(("•", "-", "*"))


LOADERS = {
    "pdf": load_pdf,
    "html": load_html,
    "nxml": load_nxml,
    "docx": load_docx,
    "markdown": load_markdown,
}
