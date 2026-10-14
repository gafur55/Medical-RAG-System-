"""Tests for src/loaders.py. Every fixture document is built here, in
tmp_path or in memory - nothing is downloaded and nothing in data/ is read."""

import re

import pytest

from src import chunking, config, loaders
from src.chunking import Unit
from src.sources import parse_registry


def kinds(units):
    return [u.kind for u in units]


def texts(units, kind=None):
    return [u.text for u in units if kind is None or u.kind == kind]


def find(units, needle):
    """The one unit containing needle."""
    hits = [u for u in units if needle in u.text]
    assert len(hits) == 1, f"{needle!r} found in {len(hits)} units"
    return hits[0]


def source(tmp_path, filename, fmt, **over):
    entry = {"id": "demo", "title": "Demo document", "type": "review", "domain": "clinical",
             "format": fmt, "path": filename, "license": "public domain"}
    entry.update(over)
    (src,) = parse_registry({"sources": [entry]}, root=tmp_path)
    return src


# ============================================================ helpers
def test_heading_key_ignores_case_numbering_and_dates():
    assert loaders.heading_key("Latest Updates to This Summary (05/12/2025)") == \
        "latest updates to this summary"
    assert loaders.heading_key("F.  References:") == "references"
    assert loaders.heading_key("1.2 Review Questions") == "review questions"
    assert loaders.heading_key("A Guide to Mohs") == "a guide to mohs"


def test_push_heading_replaces_same_or_deeper_levels():
    path = loaders.push_heading((), 1, "A")
    path = loaders.push_heading(path, 3, "A.1")
    path = loaders.push_heading(path, 2, "B")
    assert path == ((1, "A"), (2, "B"))


def test_parse_options_defaults_and_errors():
    assert loaders.parse_options({}) == {
        "skip_tables": False, "exclude_sections": loaders.DEFAULT_EXCLUDE_SECTIONS}
    assert loaders.parse_options({"exclude_sections": []})["exclude_sections"] == ()
    with pytest.raises(ValueError) as err:
        loaders.parse_options({"skip_table": True}, "billing")
    msg = str(err.value)
    assert "billing" in msg and "skip_table" in msg
    assert "skip_tables" in msg and "exclude_sections" in msg     # lists the allowed keys
    with pytest.raises(ValueError):
        loaders.parse_options({"skip_tables": "yes"})
    with pytest.raises(ValueError):
        loaders.parse_options({"exclude_sections": "References"})


def test_table_texts_split_rows_and_repeat_header():
    rows = [f"1731{i} | Mohs stage {i} with a fairly long descriptor text" for i in range(10)]
    groups = loaders.table_texts(rows, header=["Code | Descriptor"], caption="Table 1. Codes",
                                 limit=200)
    assert len(groups) > 2
    assert all(len(g) <= 200 for g in groups)
    for g in groups:
        assert g.startswith("Table 1. Codes\nCode | Descriptor\n")
    body = [line for g in groups for line in g.split("\n")[2:]]
    assert body == rows                          # every row once, in order
    assert loaders.table_texts([], header=["A | B"]) == ["A | B"]
    assert loaders.table_texts([]) == []


def test_row_text_drops_trailing_empty_cells():
    assert loaders.row_text(["a", " b\n c ", "", ""]) == "a | b c"
    assert loaders.row_text(["", " "]) == ""


def test_join_lines_mends_line_breaks():
    words = {"ill-defined", "reconstruction"}
    assert loaders.join_lines(["flap recon-", "struction is"], words) == "flap reconstruction is"
    assert loaders.join_lines(["an unknown hyphen-", "ation"]) == "an unknown hyphenation"
    assert loaders.join_lines(["complex or ill-", "defined tumor"], words) == \
        "complex or ill-defined tumor"
    assert loaders.join_lines(["non-", "melanoma skin cancer"]) == "non-melanoma skin cancer"
    assert loaders.join_lines(["codes 11102-", "11107 apply"]) == "codes 11102-11107 apply"
    assert loaders.join_lines(["Procedure-to-", "Procedure edits"]) == \
        "Procedure-to-Procedure edits"
    assert loaders.join_lines(["pre-", "and post-operative"]) == "pre- and post-operative"


def test_caption_regex_matches_ingest():
    """Read from the file, not imported: ingest.py imports the embedder."""
    source_text = (config.ROOT / "src" / "ingest.py").read_text()
    m = re.search(r'CAPTION_RE = re\.compile\(r"([^"]+)"', source_text)
    if not m:
        pytest.skip("ingest.CAPTION_RE not found")
    assert loaders.CAPTION_RE.pattern == m.group(1)


# ============================================================== load()
def test_load_rejects_the_textbook_format(tmp_path):
    (tmp_path / "book.pdf").write_bytes(b"%PDF")
    with pytest.raises(ValueError, match="pdf_textbook"):
        loaders.load(source(tmp_path, "book.pdf", "pdf_textbook", type="textbook"))


def test_load_rejects_unknown_options(tmp_path):
    (tmp_path / "a.md").write_text("# A\n\ntext\n")
    with pytest.raises(ValueError, match="allowed|Allowed"):
        loaders.load(source(tmp_path, "a.md", "markdown", options={"skip_tabels": True}))


def test_load_dispatches_on_format_and_passes_options(tmp_path):
    (tmp_path / "a.md").write_text(
        "# Protocol\n\nConsent is signed before the first stage.\n\n"
        "| Item | Owner |\n|---|---|\n| Consent | Nurse |\n")
    units = loaders.load(source(tmp_path, "a.md", "markdown"))
    assert kinds(units) == ["prose", "table"]
    units = loaders.load(source(tmp_path, "a.md", "markdown", options={"skip_tables": True}))
    assert kinds(units) == ["prose"]


# ================================================================ HTML
BOOKSHELF = """<!DOCTYPE html>
<html><head><title>Mohs Micrographic Surgery - StatPearls - NCBI Bookshelf</title>
<script>var tracking = "SCRIPT TEXT";</script><style>.x{color:red}</style></head>
<body>
<header>NCBI header SITE HEADER</header>
<nav class="toolbar">Search NCBI NAV LINKS</nav>
<div id="maincontent">
 <div class="main-content">
  <h1>Mohs Micrographic Surgery</h1>
  <div class="bk_prnt">Print view PRINT WIDGET</div>
  <h2>Continuing Education Activity</h2>
  <p>This activity reviews the procedure. OBJECTIVES TEXT</p>
  <ul><li>Identify the indications. OBJECTIVE ITEM</li></ul>
  <h2>Introduction</h2>
  <p>Mohs micrographic surgery examines <b>100%</b> of the margin.<br>It was described
     by Frederic Mohs.[1][2]</p>
  <p>The technique spares tissue.<figure><img src="x.jpg" alt="ALT TEXT">
     <a class="enlarge">Enlarge</a><figcaption>Figure 1. Horizontal frozen sections.</figcaption>
     </figure></p>
  <h3>History</h3>
  <ul><li>Chemosurgery with zinc chloride paste.<ul><li>Fixed-tissue technique.</li></ul></li>
      <li>Fresh-tissue technique.</li></ul>
  <div class="iconblock whole_rhythm fig" id="fig1">
    <a href="/figure/1"><img src="y.jpg"></a>
    <div class="icnblk_cntnt"><h4>Figure</h4>
      <p class="float-caption">Mohs map of a nasal defect. Contributed by the authors.</p></div>
  </div>
  <h2>Indications</h2>
  <table>
   <caption>Table 1. Appropriate use<sup>a</sup></caption>
   <thead><tr><th>Site</th><th>Tumor</th></tr></thead>
   <tfoot><tr><td colspan="2"><sup>a</sup>Adapted from the 2012 AUC.</td></tr></tfoot>
   <tbody>
    <tr><td>Area H</td><td>Basal cell carcinoma, any size</td></tr>
    <tr><td>Area M</td><td>Squamous cell carcinoma over 1 cm</td></tr>
   </tbody>
  </table>
  <h2>Review Questions</h2>
  <ul><li><a href="q">Access free multiple choice questions on this topic.</a></li>
      <li><a href="c">Comment on this article.</a></li></ul>
  <h2>References</h2>
  <dl><dt>1.</dt><dd><div class="bk_ref">Mohs FE. Chemosurgery: a microscopically controlled
      method of cancer excision. Arch Surg. 1941;42:279-295. [PubMed: 12345]</div></dd>
      <dt>2.</dt><dd><div class="bk_ref">Shriner DL, et al. Mohs micrographic surgery.
      J Am Acad Dermatol. 1998;39:79-97. [PMC free article: PMC999]</div></dd></dl>
  <div class="post-content"><div>Disclosure: DISCLOSURE TEXT</div></div>
 </div>
</div>
<aside><h3>Similar articles</h3><p>SIMILAR ARTICLE</p></aside>
<div id="rightcolumn"><div class="portlet">Recent Activity RECENT</div></div>
<footer>FOOTER TEXT</footer>
</body></html>"""


def test_html_sections_kinds_and_locators():
    units = loaders.html_units(BOOKSHELF, title="StatPearls: Mohs")
    intro = find(units, "examines 100% of the margin")
    assert intro.section == ("Mohs Micrographic Surgery", "Introduction")
    assert intro.locator == "Introduction" and intro.page is None
    assert intro.kind == "prose"
    assert "margin. It was described by Frederic Mohs.[1][2]" in intro.text   # <br> -> space
    history = find(units, "zinc chloride")
    assert history.section == ("Mohs Micrographic Surgery", "Introduction", "History")
    assert history.locator == "History"


def test_html_strips_scripts_navigation_and_page_furniture():
    joined = " ".join(texts(loaders.html_units(BOOKSHELF)))
    for junk in ("SCRIPT TEXT", "SITE HEADER", "NAV LINKS", "PRINT WIDGET", "SIMILAR ARTICLE",
                 "RECENT", "FOOTER TEXT", "DISCLOSURE TEXT", "ALT TEXT", "Enlarge"):
        assert junk not in joined


def test_html_default_exclusions():
    joined = " ".join(texts(loaders.html_units(BOOKSHELF)))
    assert "OBJECTIVES TEXT" not in joined and "OBJECTIVE ITEM" not in joined
    assert "multiple choice" not in joined and "Comment on this article" not in joined


def test_html_exclude_sections_replaces_the_defaults():
    units = loaders.html_units(BOOKSHELF, exclude_sections=["indications"])
    joined = " ".join(texts(units))
    assert "OBJECTIVES TEXT" in joined                  # defaults no longer apply
    assert "Area H" not in joined                        # the named section is gone
    everything = " ".join(texts(loaders.html_units(BOOKSHELF, exclude_sections=[])))
    assert "Area H" in everything and "multiple choice" in everything


def test_html_nested_blocks_are_not_counted_twice():
    units = loaders.html_units(BOOKSHELF)
    spares = find(units, "The technique spares tissue.")
    assert spares.kind == "prose" and "Horizontal frozen" not in spares.text
    caption = find(units, "Horizontal frozen sections")
    assert caption.kind == "caption" and caption.text == "Figure 1. Horizontal frozen sections."
    parent = find(units, "zinc chloride")
    assert "Fixed-tissue" not in parent.text
    find(units, "Fixed-tissue technique.")                # the sub-item is its own unit


def test_html_bookshelf_figure_block_is_one_caption():
    caption = find(loaders.html_units(BOOKSHELF), "Mohs map of a nasal defect")
    assert caption.kind == "caption"
    assert caption.text.startswith("Figure Mohs map")
    assert caption.section[-1] == "History"               # its <h4> is not a heading


def test_html_table_caption_header_and_footnote():
    (table,) = [u for u in loaders.html_units(BOOKSHELF) if u.kind == "table"]
    assert table.text.split("\n") == [
        "Table 1. Appropriate use",                       # footnote mark "a" dropped
        "Site | Tumor",
        "Area H | Basal cell carcinoma, any size",
        "Area M | Squamous cell carcinoma over 1 cm",
        "Adapted from the 2012 AUC.",
    ]
    assert table.section == ("Mohs Micrographic Surgery", "Indications")
    assert loaders.html_units(BOOKSHELF, skip_tables=True) == [
        u for u in loaders.html_units(BOOKSHELF) if u.kind != "table"]


def test_html_long_table_is_split_into_row_groups(monkeypatch):
    monkeypatch.setattr(config, "CHUNK_TARGET_CHARS", 160)
    rows = "".join(f"<tr><td>1731{i}</td><td>Mohs stage {i}, each additional block</td></tr>"
                   for i in range(8))
    html = (f"<main><h2>Codes</h2><table><tr><th>Code</th><th>Descriptor</th></tr>"
            f"{rows}</table></main>")
    tables = [u for u in loaders.html_units(html) if u.kind == "table"]
    assert len(tables) > 1
    assert all(len(t.text) <= 160 for t in tables)
    assert all(t.text.startswith("Code | Descriptor\n") for t in tables)
    assert sum(t.text.count("Mohs stage") for t in tables) == 8


def test_html_references_one_unit_per_entry():
    refs = [u for u in loaders.html_units(BOOKSHELF) if u.kind == "reference"]
    assert [r.text for r in refs] == [
        "1. Mohs FE. Chemosurgery: a microscopically controlled method of cancer excision. "
        "Arch Surg. 1941;42:279-295.",
        "2. Shriner DL, et al. Mohs micrographic surgery. J Am Acad Dermatol. 1998;39:79-97.",
    ]
    assert all(r.section[-1] == "References" and r.locator == "References" for r in refs)


NCI = """<html><body><main><article>
<h1>Skin Cancer Treatment (PDQ)</h1>
<div class="pdq-hp-patient-toggle"><a href="/patient">Go to Patient Version</a></div>
<nav class="on-this-page"><h6>On This Page</h6><a href="#a">General Information</a></nav>
<section id="_1"><h2>General Information</h2>
 <div class="pdq-sections">
  <nav class="in-this-section"><h6>In This Section</h6><ul><li>Incidence</li></ul></nav>
  <p>Skin cancer is the most common malignancy diagnosed in the United States.[1]</p>
  <section id="_2"><h3>Incidence</h3><p>BCC is more common than SCC.[2]</p></section>
  <section id="_3"><h3>Current Clinical Trials</h3><p>Use our clinical trial search.</p></section>
  <h6>References</h6>
  <ol><li>Rogers HW: Incidence estimate. Arch Dermatol 146: 283-7, 2010.
          <a href="pubmed">[PUBMED Abstract]</a></li>
      <li>Koh HK: Cutaneous melanoma. N Engl J Med 325: 171-82, 1991.</li></ol>
 </div>
</section>
<section id="_4"><h2>Latest Updates to This Summary (05/12/2025)</h2>
 <p>Editorial changes were made to this summary.</p></section>
<section id="_5"><h2>About This PDQ Summary</h2><section><h3>Purpose</h3>
 <p>This PDQ summary is for health professionals.</p></section></section>
</article></main>
<footer class="article-footer">Updated: May 12, 2025</footer></body></html>"""


def test_html_sections_scope_their_headings():
    """NCI's <h6>References follows two <h3> subsections inside the <h2>
    section. It belongs to the <h2>, not to the excluded "Current Clinical
    Trials" <h3> - without scoping, the references would vanish with it."""
    units = loaders.html_units(NCI)
    refs = [u for u in units if u.kind == "reference"]
    assert [r.text for r in refs] == [
        "1. Rogers HW: Incidence estimate. Arch Dermatol 146: 283-7, 2010.",
        "2. Koh HK: Cutaneous melanoma. N Engl J Med 325: 171-82, 1991.",
    ]
    assert refs[0].section == ("Skin Cancer Treatment (PDQ)", "General Information", "References")
    assert find(units, "most common malignancy").section == (
        "Skin Cancer Treatment (PDQ)", "General Information")


def test_html_nci_boilerplate_is_dropped():
    joined = " ".join(texts(loaders.html_units(NCI)))
    for junk in ("Patient Version", "On This Page", "In This Section", "clinical trial search",
                 "Editorial changes", "for health professionals", "Updated: May"):
        assert junk not in joined
    assert "BCC is more common than SCC.[2]" in joined


def test_html_table_class_without_a_table_is_walked_as_content():
    html = ('<main><h2>Setup</h2><div class="table"><h3>Room layout</h3>'
            '<p>The cryostat sits next to the staining station.</p></div></main>')
    units = loaders.html_units(html)
    assert [(u.text, u.kind, u.section) for u in units] == [
        ("The cryostat sits next to the staining station.", "prose", ("Setup", "Room layout"))]


def test_html_before_any_heading_uses_the_title_as_locator():
    units = loaders.html_units("<body><p>Intro text before any heading.</p><h2>A</h2>"
                               "<p>Under A.</p></body>", title="Demo doc")
    assert (units[0].section, units[0].locator) == ((), "Demo doc")
    assert (units[1].section, units[1].locator) == (("A",), "A")


def test_html_repeated_text_under_one_locator_is_kept_once():
    """ingest refuses duplicate chunk IDs; the same reference listed under
    two "References" headings would be one."""
    html = ("<main><h2>BCC</h2><h3>References</h3><ul><li>Same citation, 2001.</li></ul>"
            "<h2>SCC</h2><h3>References</h3><ul><li>Same citation, 2001.</li></ul></main>")
    assert texts(loaders.html_units(html), "reference") == ["Same citation, 2001."]


# ================================================================ nxml
JATS = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE article PUBLIC "-//NLM//DTD JATS (Z39.96) Journal Archiving DTD v1.2 20190208//EN"
  "JATS-archivearticle1.dtd">
<article xmlns:xlink="http://www.w3.org/1999/xlink" xmlns:mml="http://www.w3.org/1998/Math/MathML"
         article-type="review-article">
 <front><article-meta>
  <title-group><article-title>Mohs&nbsp;Surgery Review</article-title></title-group>
  <abstract><p>Mohs surgery has the lowest recurrence rates.</p></abstract>
 </article-meta></front>
 <body>
  <p>Opening paragraph before any section.</p>
  <sec><title>Technique</title>
   <p>Tissue is excised at a 45&deg; bevel <xref ref-type="bibr" rid="r1">[1]</xref>.
    <fig id="f1"><label>Figure 1</label><caption><title>Bevel.</title>
      <p>The scalpel is angled.</p></caption><graphic xlink:href="f1.jpg"/></fig>
   </p>
   <sec><title>Frozen sections</title>
    <list list-type="bullet"><list-item><p>Embed the tissue.</p></list-item>
     <list-item><p>Cut 4 micron sections.</p></list-item></list>
    <table-wrap id="t1"><label>Table 1</label><caption><p>Stains</p></caption>
     <table><thead><tr><th>Stain</th><th>Use</th></tr></thead>
      <tbody><tr><td>H&amp;E</td><td>Routine</td></tr><tr><td>Toluidine blue</td><td>BCC</td></tr></tbody>
     </table>
     <table-wrap-foot><fn><p>Fresh frozen tissue.</p></fn></table-wrap-foot></table-wrap>
   </sec>
  </sec>
  <sec><title>Acknowledgments</title><p>We thank the nurses.</p></sec>
 </body>
 <back><ref-list><title>References</title>
  <ref id="r1"><label>1</label><mixed-citation publication-type="journal"><person-group>
   <name><surname>Mohs</surname><given-names>FE</given-names></name></person-group>.
   <article-title>Chemosurgery</article-title>. <source>Arch Surg</source>.
   <year>1941</year>;<volume>42</volume>:<fpage>279</fpage>-<lpage>95</lpage>.
   <pub-id pub-id-type="pmid">999</pub-id></mixed-citation></ref>
  <ref id="r2"><label>2</label><element-citation publication-type="journal"><person-group
   person-group-type="author"><name><surname>Rowe</surname><given-names>DE</given-names></name>
   <name><surname>Carroll</surname><given-names>RJ</given-names></name></person-group>
   <article-title>Long-term recurrence rates</article-title><source>J Dermatol Surg Oncol</source>
   <year>1989</year><volume>15</volume><issue>3</issue><fpage>315</fpage><lpage>28</lpage>
  </element-citation></ref>
 </ref-list></back>
</article>"""


def test_nxml_sections_and_kinds(tmp_path):
    path = tmp_path / "a.nxml"
    path.write_text(JATS)
    units = loaders.load(source(tmp_path, "a.nxml", "nxml"))
    title = "Mohs Surgery Review"                   # &nbsp; resolved, then collapsed
    assert find(units, "lowest recurrence").section == (title, "Abstract")
    assert find(units, "Opening paragraph").section == (title,)
    bevel = find(units, "45° bevel")
    assert bevel.section == (title, "Technique") and bevel.locator == "Technique"
    assert bevel.text == "Tissue is excised at a 45° bevel [1]."
    caption = find(units, "scalpel is angled")
    assert caption.kind == "caption" and caption.text == "Figure 1 Bevel. The scalpel is angled."
    embed = find(units, "Embed the tissue.")
    assert embed.section == (title, "Technique", "Frozen sections")
    (table,) = [u for u in units if u.kind == "table"]
    assert table.text.split("\n") == ["Table 1 Stains", "Stain | Use", "H&E | Routine",
                                      "Toluidine blue | BCC", "Fresh frozen tissue."]
    assert "We thank the nurses" not in " ".join(texts(units))        # default exclusion


def test_nxml_references_render_both_citation_styles(tmp_path):
    units = loaders.nxml_units(JATS)
    refs = [u for u in units if u.kind == "reference"]
    assert [r.text for r in refs] == [
        "1. Mohs FE. Chemosurgery. Arch Surg. 1941;42:279-95.",
        "2. Rowe DE, Carroll RJ. Long-term recurrence rates. J Dermatol Surg Oncol. "
        "1989;15(3):315-28.",
    ]
    assert refs[0].section == ("Mohs Surgery Review", "References")


def test_nxml_default_namespace_and_skip_tables():
    xml = ('<book-part-wrapper xmlns="http://jats.nlm.nih.gov/ns/bits"><book-part><book-part-meta>'
           '<title-group><title>StatPearls Chapter</title></title-group></book-part-meta><body>'
           '<sec><title>Introduction</title><p>Mohs surgery is tissue sparing.</p>'
           '<table-wrap><table><tr><td>a</td><td>b</td></tr></table></table-wrap></sec>'
           '</body></book-part></book-part-wrapper>')
    units = loaders.nxml_units(xml, skip_tables=True)
    assert [(u.section, u.kind) for u in units] == [
        (("StatPearls Chapter", "Introduction"), "prose")]


# ================================================================ docx
def make_docx(path):
    import docx
    doc = docx.Document()
    doc.add_paragraph("Clinic Mohs Protocol", style="Title")
    doc.add_paragraph("Applies to every Mohs case.")
    doc.add_heading("Before surgery", level=1)
    doc.add_paragraph("Confirm the biopsy report and the site photograph.")
    table = doc.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Item", "Owner"), ("Consent", "Nurse"), ("Site marking", "Surgeon")]):
        table.cell(r, 0).text, table.cell(r, 1).text = a, b
    doc.add_paragraph("Table 1. Pre-operative checklist", style="Caption")
    doc.add_heading("Anticoagulation", level=2)
    doc.add_paragraph("Do not stop warfarin without the prescriber.")
    doc.add_heading("References", level=1)
    doc.add_paragraph("Alam M, et al. Bleeding risk in Mohs surgery. Dermatol Surg. 2013.")
    doc.save(path)


def test_docx_reads_body_in_order(tmp_path):
    make_docx(tmp_path / "p.docx")
    units = loaders.load(source(tmp_path, "p.docx", "docx"))
    assert kinds(units) == ["prose", "prose", "table", "caption", "prose", "reference"]
    assert units[0].section == ("Clinic Mohs Protocol",)
    assert units[0].locator == "Clinic Mohs Protocol"
    assert units[1].section == ("Clinic Mohs Protocol", "Before surgery")
    assert units[2].text == "Item | Owner\nConsent | Nurse\nSite marking | Surgeon"
    assert units[4].section == ("Clinic Mohs Protocol", "Before surgery", "Anticoagulation")
    assert units[5].section == ("Clinic Mohs Protocol", "References")
    assert all(u.page is None for u in units)


def test_docx_skip_tables_and_heading_levels(tmp_path):
    make_docx(tmp_path / "p.docx")
    units = loaders.load_docx(tmp_path / "p.docx", skip_tables=True)
    assert "table" not in kinds(units)
    assert loaders.docx_heading_level("Heading 3") == 3
    assert loaders.docx_heading_level("Title") == 0
    assert loaders.docx_heading_level("Normal") is None


# ============================================================ markdown
MARKDOWN = """---
author: clinic
---
# Mohs day protocol

Patients arrive at 7:30. They are told
to eat breakfast.

## Checklist
- Consent signed
- Site photographed
  and confirmed
- Anticoagulants noted

| Step | Owner |
|:-----|------:|
| Prep | Nurse |
| Map \\| ink | Surgeon |

```
code_line_one()
```

Wound care
----------
Keep the dressing dry for **48 hours**. See [the leaflet](https://x.org/leaflet).

## References
- Smith J. Mohs logistics. 2020.
- Jones K. Same-day repair. 2021.
"""


def test_markdown_sections_lists_tables_code():
    units = loaders.markdown_units(MARKDOWN, title="Day protocol")
    assert units[0].text == "Patients arrive at 7:30. They are told to eat breakfast."
    assert units[0].section == ("Mohs day protocol",)
    checklist = find(units, "Consent signed")
    assert checklist.text == ("- Consent signed - Site photographed and confirmed "
                              "- Anticoagulants noted")              # one paragraph
    assert checklist.section == ("Mohs day protocol", "Checklist")
    table = find(units, "Step | Owner")
    assert table.kind == "table"
    assert table.text == "Step | Owner\nPrep | Nurse\nMap | ink | Surgeon"
    code = find(units, "code_line_one")
    assert code.kind == "prose" and "```" not in code.text
    wound = find(units, "dressing dry")
    assert wound.section == ("Mohs day protocol", "Wound care")        # setext heading
    assert wound.text == "Keep the dressing dry for 48 hours. See the leaflet."
    refs = [u for u in units if u.kind == "reference"]
    assert [r.text for r in refs] == ["Smith J. Mohs logistics. 2020.",
                                      "Jones K. Same-day repair. 2021."]
    assert "author: clinic" not in " ".join(texts(units))            # front matter


def test_plain_text_without_headings(tmp_path):
    (tmp_path / "notes.txt").write_text("First paragraph of notes.\n\nSecond one here.\n")
    units = loaders.load(source(tmp_path, "notes.txt", "markdown", title="Clinic notes"))
    assert [(u.text, u.section, u.locator, u.page) for u in units] == [
        ("First paragraph of notes.", (), "Clinic notes", None),
        ("Second one here.", (), "Clinic notes", None)]


def test_markdown_exclude_sections_drops_subsections():
    md = "# Doc\n\n## Funding\n\nGrant 1.\n\n### Details\n\nGrant detail.\n\n## Results\n\nKept.\n"
    assert texts(loaders.markdown_units(md)) == ["Kept."]


# ================================================================= PDF
BODY = 11


def make_pdf(path):
    """3 pages: a running header and a page-number footer on each, a large
    heading, a bold body-size heading, a hyphenated line break, a paragraph
    running on from page 2 to page 3, a caption, a TOC line, and references."""
    import pymupdf
    doc = pymupdf.open()
    pages = [
        [("Mohs Billing Guide", 18, "hebo"),
         None,
         ("Mohs surgery is billed per stage. The first stage of the", BODY, "helv"),
         ("tumor is code 17311 and the recon-", BODY, "helv"),
         ("struction is billed separately.", BODY, "helv"),
         None,
         ("Figure 1. Stage diagram for billing.", BODY, "helv"),
         None,
         ("Contents ......................... 3", BODY, "helv")],
        [("Repairs", BODY, "hebo"),
         None,
         ("Intermediate repairs are reported in addition to the Mohs", BODY, "helv"),
         ("codes when they are medically necessary and the", BODY, "helv")],
        [("documentation supports the repair length.", BODY, "helv"),
         None,
         ("Complex repairs need undermining.", BODY, "helv"),
         None,
         ("References", BODY, "hebo"),
         None,
         ("1. Smith J. Billing for Mohs surgery. J Derm. 2020;1:1-5.", BODY, "helv"),
         ("2. Jones K. Repairs after Mohs. J Derm. 2021;2:6-9.", BODY, "helv")],
    ]
    for number, lines in enumerate(pages, start=1):
        page = doc.new_page(width=612, height=792)
        page.insert_text((72, 40), "Clinic Policy Manual - Revised 2026", fontsize=9, fontname="helv")
        page.insert_text((280, 770), f"Page {number} of 3", fontsize=9, fontname="helv")
        y = 100
        for line in lines:
            if line is None:
                y += 12           # a blank line's worth of space: a new paragraph
                continue
            text, size, font = line
            page.insert_text((72, y), text, fontsize=size, fontname=font)
            y += size * 1.3
    doc.save(path)
    doc.close()


@pytest.fixture
def pdf_units(tmp_path):
    make_pdf(tmp_path / "guide.pdf")
    return loaders.load(source(tmp_path, "guide.pdf", "pdf", type="policy"))


def test_pdf_drops_running_header_footer_and_toc(pdf_units):
    joined = " ".join(texts(pdf_units))
    assert "Clinic Policy Manual" not in joined
    assert "Page 1 of 3" not in joined and "of 3" not in joined
    assert "Contents" not in joined


def test_pdf_headings_pages_and_locators(pdf_units):
    first = find(pdf_units, "billed per stage")
    assert first.section == ("Mohs Billing Guide",)
    assert (first.page, first.locator, first.kind) == (1, "p. 1", "prose")
    repairs = find(pdf_units, "Intermediate repairs")
    assert repairs.section == ("Mohs Billing Guide", "Repairs")    # bold body size = level 2
    assert (repairs.page, repairs.locator) == (2, "p. 2")
    complex_ = find(pdf_units, "Complex repairs")
    assert (complex_.page, complex_.section) == (3, ("Mohs Billing Guide", "Repairs"))


def test_pdf_mends_hyphenation_and_page_breaks(pdf_units):
    first = find(pdf_units, "billed per stage")
    assert "the tumor is code 17311 and the reconstruction is billed separately." in first.text
    # A sentence cut by the page break stays one paragraph, cited at its start.
    repairs = find(pdf_units, "Intermediate repairs")
    assert repairs.text.endswith("medically necessary and the documentation supports the "
                                 "repair length.")
    assert repairs.page == 2


def test_pdf_captions_and_references(pdf_units):
    assert find(pdf_units, "Stage diagram").kind == "caption"
    refs = [u for u in pdf_units if u.kind == "reference"]
    assert [r.text for r in refs] == [
        "1. Smith J. Billing for Mohs surgery. J Derm. 2020;1:1-5.",
        "2. Jones K. Repairs after Mohs. J Derm. 2021;2:6-9."]
    assert refs[0].section == ("Mohs Billing Guide", "References")
    assert refs[0].locator == "p. 3"


def test_pdf_exclude_sections(tmp_path):
    make_pdf(tmp_path / "guide.pdf")
    units = loaders.load_pdf(tmp_path / "guide.pdf", exclude_sections=["repairs"])
    joined = " ".join(texts(units))
    assert "Intermediate repairs" not in joined and "Complex repairs" not in joined
    assert "billed per stage" in joined


def test_pdf_exclusion_stops_at_the_next_heading_of_any_level(tmp_path):
    """PDF heading levels are guesses. Excluding the 18pt "Mohs Billing
    Guide" heading drops its own text, not the bold "Repairs" section that
    ranks below it."""
    make_pdf(tmp_path / "guide.pdf")
    units = loaders.load_pdf(tmp_path / "guide.pdf", exclude_sections=["mohs billing guide"])
    joined = " ".join(texts(units))
    assert "billed per stage" not in joined
    assert "Intermediate repairs" in joined


def test_heading_levels_rank_by_size_but_not_one_off_banners():
    # 24pt once on page 1 = the title; 20pt once on page 3 = a front-matter
    # banner; 16pt chapters and 13pt sections recur.
    headings = [(24, 1), (20, 3), (16, 5), (16, 40), (13, 6), (13, 41)]
    assert loaders.heading_levels(headings) == {24: 1, 20: 2, 16: 2, 13: 3}
    assert loaders.heading_levels([]) == {}


def test_overlong_reference_block_is_split(monkeypatch):
    monkeypatch.setattr(config, "CHUNK_TARGET_CHARS", 200)
    entries = " ".join(f"[{i}] Author{i} A. A study of flaps number {i}. J Derm. 2020;1:{i}."
                       for i in range(1, 9))
    units = loaders.markdown_units(f"# References\n\n{entries}\n")
    assert len(units) > 1 and all(u.kind == "reference" for u in units)
    assert all(len(u.text) <= 200 for u in units)
    assert " ".join(u.text for u in units) == entries


def test_pdf_known_heading_without_a_gap_is_split_off():
    """A bold "References" set straight above the first entry (no blank
    line, no gap) still becomes a heading, so the entries are references."""
    def line(text, y, bold=False):
        return {"text": text, "size": 10, "bold": bold, "y0": y, "y1": y + 12}
    page = [line("Results", 100, bold=True), line("Mohs cured 99% of tumors.", 125),
            line("References", 160, bold=True),
            line("[1] Mohs FE. Chemosurgery. Arch Surg. 1941;42:279.", 172),
            line("[2] Rowe DE. Recurrence rates. J Derm Surg. 1989;15:315.", 184)]
    units = loaders.pdf_units([(800, page)])
    assert [(u.kind, u.section) for u in units] == [
        ("prose", ("Results",)), ("reference", ("References",)), ("reference", ("References",))]


def test_furniture_needs_repetition_in_the_margin():
    def line(text, y):
        return {"text": text, "size": 10, "bold": False, "y0": y, "y1": y + 10}
    pages = [(800, [line("Running title", 20), line(f"Body text {i}", 400),
                    line("Once only footer", 780)]) for i in range(1)]
    pages += [(800, [line("Running title", 20), line(f"Body text {i}", 400)]) for i in range(3)]
    found = loaders.furniture_texts(pages)
    assert found == {"running title"}                     # body repeats too, but not in the margin


# ======================================================== pack() on top
def test_pack_over_loader_output_never_crosses_sections():
    """Many paragraphs per section, each tagged with its section's token;
    every packed chunk may only carry its own section's token."""
    md = []
    for s in range(4):
        md.append(f"## Section {s}\n")
        for p in range(6):
            md.append(f"Paragraph {p} of part tok{s}x. " + " ".join(
                f"Sentence {k} about tok{s}x and margins." for k in range(5)) + "\n")
    units = loaders.markdown_units("\n".join(md))
    chunks = chunking.pack(units, config.CHUNK_TARGET_CHARS, config.CHUNK_OVERLAP_CHARS,
                           config.CHUNK_MIN_CHARS)
    assert len(chunks) > 4
    for chunk in chunks:
        (s,) = {int(t) for t in re.findall(r"tok(\d)x", chunk["text"])}
        assert chunk["section"] == (f"Section {s}",)
        assert chunk["locator"] == f"Section {s}"


def test_pack_over_html_and_pdf_units_keeps_kinds_and_places(pdf_units):
    for units in (loaders.html_units(BOOKSHELF), loaders.html_units(NCI), pdf_units):
        chunks = chunking.pack(units, config.CHUNK_TARGET_CHARS, config.CHUNK_OVERLAP_CHARS,
                               config.CHUNK_MIN_CHARS)
        places = {(u.section, u.locator, u.page, u.kind) for u in units}
        assert chunks
        for c in chunks:
            assert (c["section"], c["locator"], c["page"], c["kind"]) in places
        keys = [(c["page"] if c["page"] is not None else c["locator"], c["text"]) for c in chunks]
        assert len(keys) == len(set(keys))                # no duplicate chunk IDs


def test_units_are_chunking_units():
    units = loaders.markdown_units("# A\n\nSome text here.\n")
    assert units == [Unit("Some text here.", ("A",), "A", None, "prose")]
